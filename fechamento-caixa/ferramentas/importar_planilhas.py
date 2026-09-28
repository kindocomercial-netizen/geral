"""Extrai os fechamentos diários das planilhas "Backup de Fecha caixa" (.xlk/.xlsx/.xlsm)
e gera um JSON no formato de backup do app (Configurações > Restaurar backup).

Uso: python3 importar_planilhas.py arquivo1.xlk [arquivo2.xlk ...] -o historico.json
     python3 importar_planilhas.py arquivo1.xlk [...] -o ../dados.js   (o app carrega sozinho na 1ª abertura)
"""
import argparse, io, json, re, sys, unicodedata
from datetime import date
import openpyxl

MESES = {'janeiro': 1, 'fevereiro': 2, 'marco': 3, 'abril': 4, 'maio': 5, 'junho': 6, 'julho': 7,
         'agosto': 8, 'setembro': 9, 'outubro': 10, 'novembro': 11, 'dezembro': 12}


def normalizar(s):
    s = unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'\s+', ' ', s).strip()


def centavos(v):
    if isinstance(v, (int, float)):
        return round(v * 100)
    return 0


def campo(rotulo):
    r = normalizar(rotulo)
    if r.startswith('rel. caixa') or r.startswith('rel caixa'): return 'sistema'
    if r.startswith('troco inic'): return 'trocoInicial'
    if 'suprimento' in r: return 'suprimento'
    if 'cred a vista' in r or 'pos cred' in r: return 'credito'
    if 'cred parcelado' in r: return 'creditoParcelado'
    if r.startswith('debito'): return 'debito'
    if r.startswith('pix'): return 'pix'
    if r.startswith('troco final'): return 'trocoFinal'
    if r.startswith('sangria'): return 'sangrias'
    if r.startswith('despesa'): return 'despesas'
    if r.startswith('vale transp'): return 'valeTransporte'
    return None


def mes_do_arquivo(nome):
    n = normalizar(nome)
    for m, i in MESES.items():
        if m in n:
            ano = re.search(r'20\d\d', n)
            return i, int(ano.group()) if ano else None
    return None, None


def data_da_aba(titulo, mes_arq, ano_arq):
    m = re.fullmatch(r'(\d{1,2})_(\d{1,2})_(\d{4})', titulo.strip())
    if not m:
        return None, None
    d, mes, ano = map(int, m.groups())
    aviso = None
    if mes_arq and mes != mes_arq:
        aviso = f'aba "{titulo}" está no arquivo de {mes_arq:02d}/{ano_arq}; considerada {d:02d}/{mes_arq:02d}/{ano_arq}'
        mes, ano = mes_arq, ano_arq or ano
    return date(ano, mes, d).isoformat(), aviso


def ler_aba(ws):
    """Localiza cada 'Caixa Nº X' na linha 1 e lê pares rótulo/valor abaixo dele."""
    caixas = []
    for c in ws[1]:
        if c.value and normalizar(c.value).startswith('caixa'):
            num = re.search(r'\d+', str(c.value))
            caixas.append((int(num.group()) if num else len(caixas) + 1, c.column))
    resultado = []
    for numero, col in caixas:
        cx = {'caixa': numero, 'sangrias': [], 'outros': []}
        na_gaveta = False
        for linha in range(2, 30):
            rot = ws.cell(linha, col).value
            if not isinstance(rot, str):
                continue
            r = normalizar(rot)
            if r.startswith('total sistema'):
                na_gaveta = True
                continue
            if r.startswith('total gaveta'):
                break
            valor = centavos(ws.cell(linha, col + 1).value)
            chave = campo(rot)
            if chave == 'sangrias':
                if valor:
                    cx['sangrias'].append(valor)
            elif chave:
                cx[chave] = cx.get(chave, 0) + valor
            elif na_gaveta and valor:
                # Linha com rótulo fora do padrão (ex.: "CED", "crediario"): a planilha soma na gaveta.
                cx['outros'].append({'desc': rot.strip(), 'valor': valor})
        if not cx['outros']:
            del cx['outros']
        resultado.append(cx)
    return resultado


def carregar(caminho):
    with open(caminho, 'rb') as fh:
        return openpyxl.load_workbook(io.BytesIO(fh.read()), data_only=True)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('arquivos', nargs='+')
    ap.add_argument('-o', '--saida', default='historico.json')
    args = ap.parse_args()

    dias, avisos = {}, []
    for caminho in args.arquivos:
        mes_arq, ano_arq = mes_do_arquivo(caminho)
        wb = carregar(caminho)
        for ws in wb.worksheets:
            iso, aviso = data_da_aba(ws.title, mes_arq, ano_arq)
            if not iso:
                continue
            if aviso:
                avisos.append(aviso)
            caixas = ler_aba(ws)
            if not caixas:
                continue
            if not any(v for c in caixas for k, v in c.items() if k != 'caixa'):
                avisos.append(f'aba "{ws.title}" de {caminho.split("/")[-1]} está em branco; ignorada')
                continue
            if iso in dias:
                avisos.append(f'{iso} aparece em mais de um arquivo; mantida a versão de {caminho}')
            dias[iso] = {'id': 'imp-' + iso, 'data': iso, 'caixas': caixas, 'obs': '',
                         'origem': caminho.split('/')[-1]}

    fechamentos = [dias[k] for k in sorted(dias)]
    with open(args.saida, 'w', encoding='utf-8') as fh:
        if args.saida.endswith('.js'):
            fh.write('window.HISTORICO_INICIAL = ' + json.dumps(fechamentos, ensure_ascii=False) + ';\n')
        else:
            json.dump({'versao': 2, 'fechamentos': fechamentos}, fh, ensure_ascii=False, indent=1)
    print(f'{len(fechamentos)} dias exportados para {args.saida}')
    for a in avisos:
        print('aviso:', a, file=sys.stderr)


if __name__ == '__main__':
    main()
